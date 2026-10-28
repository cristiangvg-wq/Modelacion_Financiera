"""UI: Excel con hoja Parámetros para ISA → la barra lateral muestra su deuda y el VPS la usa."""
import os, sys, re
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
import openpyxl
import modelacion_logic as L
from playwright.sync_api import sync_playwright
URL, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)
ISA = "ISA — Interconexión Eléctrica"
L.vincular_estado_sesion({}); L.preparar_empresa(ISA)
df, _m, _s = L.obtener_datos_financieros_publicos(ISA)
ruta = os.path.abspath(f"{OUT}/isa_parametros.xlsx")
L.asegurar_plantilla_excel(ruta, forzar=True)
wb = openpyxl.load_workbook(ruta); ws, wp = wb["Datos"], wb["Parametros"]
cols = ["Ingresos", "EBITDA", "EBIT", "Depreciacion", "CAPEX", "NWC"]
for r, (anio, fila) in enumerate(df.iterrows(), start=4):
    ws.cell(r, 1, int(anio))
    for c, k in enumerate(cols, start=2):
        ws.cell(r, c, float(fila[k]))
vals = {"Deuda financiera": 40000, "Caja y equivalentes": 5000, "Moneda": "COP", "País": "Colombia"}
for row in wp.iter_rows(min_row=3, max_row=14):
    if row[0].value in vals:
        row[1].value = vals[row[0].value]
wb.save(ruta)
log = []
with sync_playwright() as pw:
    b = pw.chromium.launch(); pg = b.new_page(viewport={"width": 1440, "height": 2600})
    pg.goto(URL, wait_until="networkidle", timeout=120000); pg.wait_for_timeout(8000)
    for txt in ("Saltar tour", "Listo, empezar"):
        try:
            pg.get_by_role("button", name=txt).first.click(timeout=3000); pg.wait_for_timeout(2000)
        except Exception:
            pass
    pg.get_by_text("Parámetros del modelo").first.click(timeout=10000); pg.wait_for_timeout(2000)
    dn = pg.get_by_label("Deuda neta").first
    log.append("deuda neta antes: " + dn.input_value())
    pg.locator('input[type="file"]').first.set_input_files(ruta); pg.wait_for_timeout(15000)
    log.append("deuda neta después de cargar: " + pg.get_by_label("Deuda neta").first.input_value())
    log.append("caja después de cargar: " + pg.get_by_label("Caja").first.input_value())
    pg.get_by_role("button", name=re.compile("Generar todo")).first.click(timeout=10000); pg.wait_for_timeout(120000)
    pg.get_by_role("tab").nth(3).click(); pg.wait_for_timeout(8000)
    t = pg.locator('[role="tabpanel"]:visible').first.inner_text()
    for ln in t.split("\n"):
        if re.search(r"Valor por acción|Upside|Deuda neta|Fuente financiera", ln):
            log.append("[tab4] " + ln[:160])
    exc = pg.locator('[data-testid="stException"]').all_inner_texts()
    log.append(f"excepciones: {[e[:300] for e in exc]}")
    pg.screenshot(path=f"{OUT}/tab4.png")
    b.close()
open(f"{OUT}/log.txt", "w").write("\n".join(log)); print("\n".join(log))
