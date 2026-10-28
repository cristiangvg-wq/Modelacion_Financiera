"""Plantilla EEFF v2 y cargador de Excel: encabezados con títulos, formato ancho, hoja Parámetros y ejemplo."""
import os
import tempfile
import unittest

import openpyxl
import pandas as pd

import modelacion_logic as L

FILAS = [[2021, 9800, 3600, 2800, 800, 1500, 1000], [2022, 10600, 4000, 3100, 900, 1700, 1080],
         [2023, 11500, 4400, 3500, 900, 1800, 1150], [2024, 12300, 4800, 3850, 950, 1950, 1220],
         [2025, 13100, 5150, 4150, 1000, 2050, 1300]]


class _Base(unittest.TestCase):
    def setUp(self):
        L.vincular_estado_sesion({})
        self.dir = tempfile.mkdtemp()

    def tearDown(self):
        L.desvincular_estado_sesion()
        for n in list(L.EMPRESAS_RUNTIME):
            if n.startswith("Prueba EEFF"):
                L.EMPRESAS_RUNTIME.pop(n, None)

    def plantilla(self, nombre="p.xlsx"):
        ruta = os.path.join(self.dir, nombre)
        L.asegurar_plantilla_excel(ruta, forzar=True)
        return ruta


class TestPlantilla(_Base):
    def test_estructura_v2(self):
        wb = openpyxl.load_workbook(self.plantilla())
        self.assertEqual(wb.sheetnames, ["Datos", "Parametros", "Ejemplo", "Instrucciones"])
        self.assertEqual([c.value for c in wb["Datos"][3]][:7], ["Año", "Ingresos", "EBITDA", "EBIT", "Depreciacion", "CAPEX", "NWC"])
        self.assertIsNone(wb["Datos"]["B4"].value)  # se entrega vacía para llenar
        self.assertTrue(str(wb["Instrucciones"]["A30"].value).endswith("versión " + L.VERSION_PLANTILLA_EEFF))

    def test_plantilla_vieja_se_actualiza(self):
        ruta = os.path.join(self.dir, "vieja.xlsx")
        wb = openpyxl.Workbook()
        wb.active["A1"] = "x" * 3000
        wb.save(ruta)
        L.asegurar_plantilla_excel(ruta)
        self.assertIn("Parametros", openpyxl.load_workbook(ruta).sheetnames)

    def test_plantilla_llena_carga_con_parametros(self):
        ruta = self.plantilla()
        wb = openpyxl.load_workbook(ruta)
        ws, wp = wb["Datos"], wb["Parametros"]
        for r, fila in enumerate(FILAS, start=4):
            for c, v in enumerate(fila, start=1):
                ws.cell(r, c, v * (1000 if c > 1 else 1))  # cifras en millones
        valores = {"Unidad de las cifras": "millones", "Tasa de impuestos (%)": 33, "Deuda financiera": 4000000,
                   "Caja y equivalentes": 500000, "Interés minoritario": 200000, "Acciones en circulación (millones)": 1000,
                   "Precio por acción": 9500, "Moneda": "COP", "País": "Colombia"}
        for row in wp.iter_rows(min_row=3, max_row=14):
            if row[0].value in valores:
                row[1].value = valores[row[0].value]
        wb.save(ruta)
        df, msg = L.cargar_y_limpiar_excel(ruta, "Prueba EEFF 1", False)
        self.assertEqual(list(df.index), [2021, 2022, 2023, 2024, 2025])
        self.assertAlmostEqual(float(df.loc[2025, "Ingresos"]), 13100.0)  # millones → miles de millones
        self.assertAlmostEqual(float(df.loc[2025, "NOPAT"]), 4150 * (1 - 0.33), places=6)
        self.assertIn("Parámetros", msg)
        p = L.params_empresa("Prueba EEFF 1")
        self.assertAlmostEqual(p["tax"], 0.33)
        self.assertAlmostEqual(p["deuda_neta"], 3500.0, places=3)
        self.assertAlmostEqual(p["nci"], 200.0, places=3)
        self.assertAlmostEqual(p["acciones"], 1000.0, places=3)
        self.assertAlmostEqual(p["precio_referencia"], 9500.0)
        self.assertEqual(L.resolver_precio_mercado("Prueba EEFF 1"), 9500.0)
        self.assertEqual(L.EMPRESAS_RUNTIME["Prueba EEFF 1"]["moneda"], "COP")

    def test_ejemplo_sin_reemplazar_avisa(self):
        ruta = os.path.join(self.dir, "ej.xlsx")
        pd.DataFrame([[2021, 9200, 3500, 2700, 800, 1600, 1000], [2022, 10500, 4100, 3200, 900, 1800, 1100],
                      [2023, 11200, 4500, 3600, 900, 1900, 1150], [2024, 12000, 4900, 4000, 900, 2000, 1200],
                      [2025, 12900, 5300, 4350, 950, 2100, 1260]],
                     columns=["Año", "Ingresos", "EBITDA", "EBIT", "Depreciacion", "CAPEX", "NWC"]).to_excel(ruta, index=False)
        with self.assertRaisesRegex(ValueError, "ejemplo"):
            L.cargar_y_limpiar_excel(ruta, "Prueba EEFF 2", False)


class TestFormatos(_Base):
    def test_titulos_arriba_y_formato_ancho(self):
        ruta = os.path.join(self.dir, "oficial.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Estado de resultados consolidado"
        ws["A2"] = "Cifras en miles de millones de COP"
        ws.append([])
        ws.append(["Partida", 2021, 2022, 2023, 2024, 2025])
        for nombre, idx in (("Ingresos operacionales", 1), ("EBITDA", 2), ("Utilidad operativa (EBIT)", 3),
                            ("Depreciación", 4), ("CAPEX", 5), ("Capital de trabajo", 6)):
            ws.append([nombre] + [f[idx] for f in FILAS])
        wb.save(ruta)
        df, msg = L.cargar_y_limpiar_excel(ruta, "Prueba EEFF 3", False)
        self.assertEqual(len(df), 5)
        self.assertAlmostEqual(float(df.loc[2023, "CAPEX"]), 1800.0)
        self.assertIn("formato ancho", msg)

    def test_tasa_por_defecto_25(self):
        ruta = os.path.join(self.dir, "simple.csv")
        pd.DataFrame(FILAS, columns=["Año", "Ingresos", "EBITDA", "EBIT", "Depreciacion", "CAPEX", "NWC"]).to_csv(ruta, index=False)
        df, _ = L.cargar_y_limpiar_excel(ruta, "Prueba EEFF 4", False)
        self.assertAlmostEqual(float(df.loc[2025, "NOPAT"]), 4150 * 0.75, places=6)  # antes NOPAT = EBIT (0 %)

    def test_sin_tabla_mensaje_claro(self):
        ruta = os.path.join(self.dir, "vacio.xlsx")
        pd.DataFrame({"x": ["hola", "mundo"]}).to_excel(ruta, index=False)
        with self.assertRaisesRegex(ValueError, "encabezados"):
            L.cargar_y_limpiar_excel(ruta, "Prueba EEFF 5", False)


if __name__ == "__main__":
    unittest.main()
